from openpyxl import load_workbook

from analytics import financial_model


def test_model_is_formula_driven_and_complete(tmp_path, monkeypatch):
    monkeypatch.setattr(financial_model, "OUTPUT", tmp_path / "model.xlsx")
    path = financial_model.build(6421, 155.27, "2018-08-01")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Assumptions", "Model", "Summary"]
    model = wb["Model"]
    assert model["B2"].value.startswith("=Assumptions!$B$5")  # growth from the real starting point
    assert model[f"L{financial_model.MONTHS + 1}"].value.startswith("=L")  # cash rolls forward
    assert wb["Assumptions"]["B5"].value == 6421
    assert "MATCH" in wb["Assumptions"]["E9"].value  # selected-scenario lookup
